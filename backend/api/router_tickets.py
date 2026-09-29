"""API router for Ticket generation and management."""

import asyncio
import io
import json
import logging
import os
import re
import uuid
from typing import Any, Dict, List, Optional
from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, Response, UploadFile, status
from pydantic import BaseModel
from sqlalchemy import desc, func, select, text
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.core.bricks_compiler import compile_ticket
from backend.core.faker import FakeDataGenerator
from backend.core.text_normalization import expand_address_for_tts, format_phone_for_tts, expand_address
from backend.core.ticket_generator import generate_tickets
from backend.core.tts_v2 import tts_engine_v2
from backend.database import AsyncSessionLocal, get_db
from backend.models.domain_02 import GeneratedTicket, ScenarioTicket
from backend.schemas.domain_02 import (
    TicketUpdate,
    TicketFilterParams,
    TicketGenerateAcceptedResponse,
    TicketGenerateRequest,
    TicketResponse,
    TicketStatusResponse,
)
from backend.schemas.generator import ClassifierRow

logger = logging.getLogger(__name__)

api_tickets_router = APIRouter(
    prefix="/api/tickets",
    tags=["Tickets"],
)

tickets_router = APIRouter(
    prefix="/api/v1/tickets",
    tags=["Tickets"],
)

# Helper function for display_id
def make_abbr(text: str) -> str:
    if not text:
        return "ХЗ"
    return "".join([w[0].upper() for w in text.split()[:2] if w])


AVAILABLE_DEFAULT_CATEGORIES: List[str] = [
    "Пожары и задымления",
    "ДТП",
    "Запах газа",
    "Оказание медицинской скорой и неотложной помощи",
    "Нарушение правопорядка",
    "Взрывы",
    "Аварии и происшествия в городском хозяйстве",
    "Человек в опасности",
]

# Preset templates for categories
CATEGORY_TEMPLATES: Dict[str, List[str]] = {
    "дтп": [
        "Срочно! Произошло лобовое столкновение двух автомобилей на ул. {street}, рядом с домом {house}. Водитель зажат, есть пострадавшие! Я {role} {first_name}.",
        "Авария, сильное ДТП возле дома {house} по улице {street}. Машина перевернулась, требуется деблокировка и скорая! Заявитель: {last_name} {first_name}.",
        "На перекрёстке у дома {house} на улице {street} столкнулись автобус и легковушка. Много пострадавших. Звонит {role}, тел. {phone}.",
    ],
    "пожар": [
        "Здравствуйте! Горит квартира на 3 этаже, сильный дым в подъезде по адресу ул. {street}, д. {house}! Звонит {role} {first_name}. Срочно пожарных!",
        "У нас пожар! Из окна валит густой чёрный дым, улица {street}, дом {house}. Я {role}, зовут {last_name} {first_name}. Помогите!",
        "Пожар в многоквартирном доме, улица {street}, д. {house}. Огонь перекинулся на балкон. Заявитель: {role} {first_name} {middle_name}, тел. {phone}.",
    ],
    "газ": [
        "В подъезде дома {house} на улице {street} резкий сильный запах газа, режет глаза! Жильцы выходят на улицу. Я {role} {first_name} {last_name}, тел. {phone}.",
        "Утечка газа на лестничной клетке по адресу {street}, д. {house}. Боимся взрыва! Заявитель: {role} {first_name}.",
    ],
    "медицин": [
        "Скорую помощь, пожалуйста! Человеку плохо с сердцем, потерял сознание по адресу ул. {street}, д. {house}. Я {role} {first_name} {last_name}.",
        "Человек упал на улице возле дома {house}, ул. {street}, тяжело дышит и не реагирует. Звонит {role} {first_name}. Пришлите реанимацию!",
    ],
    "правопоряд": [
        "Полицию срочно! Во дворе дома {house} по ул. {street} массовая драка, бьют витрины. Я {role} {first_name} {middle_name}.",
        "Во дворе на улице {street}, около дома {house}, группа неизвестных шумит и ломает забор. Заявитель: {role} {last_name} {first_name}.",
    ],
}

DEFAULT_TEMPLATES: List[str] = [
    "Срочно требуется помощь по адресу ул. {street}, д. {house}! Происшествие: {incident_name}. Звонит {role} {first_name} {last_name}, тел. {phone}.",
    "Здравствуйте! По адресу ул. {street}, дом {house} произошло ЧП: {incident_name}! Я {role} {first_name}. Срочно пришлите помощь!",
    "Помогите, у нас тут {incident_name}! Адрес: улица {street}, дом {house}. Заявитель: {role} {last_name} {first_name}.",
]

_CLASSIFIER_EKP_DATA: Optional[Dict[str, Any]] = None


def load_classifier_ekp() -> Dict[str, Any]:
    """Lazy load data/classifier_ekp.json."""
    global _CLASSIFIER_EKP_DATA
    if _CLASSIFIER_EKP_DATA is None:
        project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        possible_paths = [
            os.path.join(project_root, "data", "classifier_ekp.json"),
            os.path.join(os.getcwd(), "data", "classifier_ekp.json"),
        ]
        for path in possible_paths:
            if os.path.exists(path):
                try:
                    with open(path, "r", encoding="utf-8") as f:
                        _CLASSIFIER_EKP_DATA = json.load(f)
                    break
                except Exception as e:
                    logger.warning("Error loading classifier_ekp.json from %s: %s", path, e)
        if _CLASSIFIER_EKP_DATA is None:
            _CLASSIFIER_EKP_DATA = {"records": []}
    return _CLASSIFIER_EKP_DATA


def get_templates_for_category(category: str, subcategory: Optional[str] = None) -> List[str]:
    combined = f"{category} {subcategory or ''}".lower()
    for key, tpls in CATEGORY_TEMPLATES.items():
        if key in combined:
            return tpls
    return DEFAULT_TEMPLATES


def extract_service_codes(record: dict, category: str) -> list:
    main_svc = record.get("main_service")
    raw_services = record.get("base_services") or record.get("services") or []
    
    # Heuristic: If there are too many services, filter to core emergency ones or just main_service
    if len(raw_services) > 4:
        core = [s for s in raw_services if any(x in str(s) for x in ["101", "102", "103", "104", "Антитеррор"])]
        if core:
            raw_services = core
        elif main_svc:
            raw_services = [main_svc]

    codes = set()
    for s in raw_services:
        s_str = str(s).strip()
        if not s_str:
            continue
        lower_s = s_str.lower()
        if "101" in lower_s:
            codes.add("01 Пожарные")
        elif "102" in lower_s:
            codes.add("02 Полиция")
        elif "103" in lower_s:
            codes.add("03 Скорая")
        elif "104" in lower_s:
            codes.add("04 Газ")
        else:
            codes.add(s_str)
            
    if not codes:
        codes.update(["01 Пожарные", "02 Полиция", "03 Скорая"])
        
    return sorted(list(codes))


def resolve_classifier_row(category: str, subcategory: Optional[str] = None) -> ClassifierRow:
    ekp_data = load_classifier_ekp()
    records = ekp_data.get("records", [])

    cat_query = (category or "").strip().lower()
    sub_query = (subcategory or "").strip().lower() if subcategory else ""

    matched_records = [
        r for r in records
        if cat_query in r.get("category", "").lower() or r.get("category", "").lower() in cat_query
    ]

    if sub_query and matched_records:
        sub_matched = [
            r for r in matched_records
            if sub_query in r.get("group", "").lower()
            or sub_query in r.get("final_type", "").lower()
            or r.get("group", "").lower() in sub_query
        ]
        if sub_matched:
            matched_records = sub_matched

    templates = get_templates_for_category(category, subcategory)

    import random
    if matched_records:
        rec = random.choice(matched_records)
        code = str(rec.get("id") or "01.01")
        incident_name = rec.get("final_type") or rec.get("group") or category
        services = extract_service_codes(rec, category)
        raw_markers = [
            rec.get("feature1"),
            rec.get("feature2"),
            rec.get("feature3"),
            rec.get("group"),
            rec.get("final_type"),
        ]
        markers = [str(m) for m in raw_markers if m]
        
        # Если подкатегория не была передана, мы назначаем её из выбранного инцидента
        if not subcategory:
            subcategory = rec.get("group") or rec.get("final_type")
    else:
        code = "01.01"
        incident_name = subcategory or category
        services = extract_service_codes({}, category)
        markers = [category]

    return ClassifierRow(
        code=code,
        incident_name=incident_name,
        services=services,
        markers=markers,
        templates=templates,
    )





active_generations: int = 0


async def generate_tickets_background_task(
    category: Optional[str],
    subcategory: Optional[str],
    count: int,
):
    """Background task to generate tickets, save to DB, and cache TTS audio for bricks."""
    global active_generations
    processed = 0
    try:
        faker = FakeDataGenerator()
        tickets_with_sub = []
        import random

        for _ in range(count):
            cur_cat = category
            if not cur_cat or cur_cat.strip().lower() in ["случайная категория", "случайная", "random", "any", "все категории"]:
                cur_cat = random.choice(AVAILABLE_DEFAULT_CATEGORIES)
            
            cur_sub = subcategory
            if not cur_sub or cur_sub.strip().lower() in ["случайная подкатегория", "случайная", "random", "any", "все подкатегории"]:
                cur_sub = None
                
            classifier_row = resolve_classifier_row(category=cur_cat, subcategory=cur_sub)
            ticket_sub = classifier_row.incident_name
            generated_list = await asyncio.to_thread(generate_tickets, classifier_row=classifier_row, count=1, faker=faker)
            for t in generated_list:
                tickets_with_sub.append((t, cur_cat, ticket_sub))

        # 1. Save to DB
        from sqlalchemy import text
        async with AsyncSessionLocal() as session:
            # Ensure sys_sequences exists and has default row
            await session.execute(
                text("CREATE TABLE IF NOT EXISTS sys_sequences (name VARCHAR PRIMARY KEY, last_val INTEGER NOT NULL DEFAULT 0)")
            )
            try:
                await session.execute(
                    text("INSERT INTO sys_sequences (name, last_val) VALUES ('generated_tickets', 0) ON CONFLICT (name) DO NOTHING")
                )
            except Exception:
                await session.execute(
                    text("INSERT OR IGNORE INTO sys_sequences (name, last_val) VALUES ('generated_tickets', 0)")
                )
            await session.commit()

            # Get persistent sequence
            seq_res = await session.execute(text("SELECT last_val FROM sys_sequences WHERE name='generated_tickets'"))
            current_seq = seq_res.scalar() or 0

            for t, cur_cat, cur_sub in tickets_with_sub:
                current_seq += 1
                gen_ticket = GeneratedTicket(
                    id=t.ticket_id,
                    category=cur_cat,
                    subcategory=cur_sub,
                    complexity=t.complexity,
                    plot=t.plot,
                    factoids=t.factoids,
                    ground_truth=t.ground_truth,
                    etalon_services=t.etalon_services,
                    sequence_number=current_seq,
                )
                session.add(gen_ticket)

                scenario_ticket = ScenarioTicket(
                    scenario_id=t.ticket_id,
                    settings={
                        "complexity": t.complexity,
                        "etalon_services": t.etalon_services,
                        "plot": t.plot,
                        "category": cur_cat,
                        "subcategory": cur_sub,
                        "sequence_number": current_seq,
                    },
                    ground_truth=t.ground_truth,
                    ai_content={
                        "factoids": t.factoids,
                        "plot": t.plot,
                    },
                    workflow_state={"status": "active"},
                )
                session.add(scenario_ticket)

            # Update persistent sequence
            await session.execute(
                text("UPDATE sys_sequences SET last_val = :val WHERE name = 'generated_tickets'"),
                {"val": current_seq}
            )
            await session.commit()

        # 2. Compile BricksMatrix and cache audio via TTS
        for t, _, _ in tickets_with_sub:
            try:
                matrix = await asyncio.to_thread(compile_ticket, t)
                for brick in matrix.bricks:
                    if brick.text:
                        try:
                            await asyncio.to_thread(tts_engine_v2.synthesize, brick.text)
                        except Exception as e:
                            logger.warning("TTS synthesize error for text '%s': %s", brick.text[:30], e)
                # Cache preview audio
                await asyncio.to_thread(_generate_preview_audio_for_ticket, t.plot or "", t.ground_truth or {})
            except Exception as e:
                logger.warning("Error compiling bricks for ticket %s: %s", t.ticket_id, e)
            finally:
                processed += 1
                active_generations = max(0, active_generations - 1)

    except asyncio.CancelledError:
        logger.warning("generate_tickets_background_task was cancelled")
    except Exception as e:
        logger.error("Error in generate_tickets_background_task: %s", e, exc_info=True)
    finally:
        remaining_to_deduct = count - processed
        if remaining_to_deduct > 0:
            active_generations = max(0, active_generations - remaining_to_deduct)



@api_tickets_router.get("/status", response_model=TicketStatusResponse)
@api_tickets_router.get("/status/", response_model=TicketStatusResponse, include_in_schema=False)
@tickets_router.get("/status", response_model=TicketStatusResponse)
@tickets_router.get("/status/", response_model=TicketStatusResponse, include_in_schema=False)
async def get_tickets_status() -> TicketStatusResponse:
    """Статус фоновых генераций для честного UI."""
    return TicketStatusResponse(
        is_generating=active_generations > 0,
        remaining_tickets=active_generations,
    )


@api_tickets_router.get("/counts")
@api_tickets_router.get("/counts/", include_in_schema=False)
@tickets_router.get("/counts")
@tickets_router.get("/counts/", include_in_schema=False)
async def get_tickets_counts(
    complexity: Optional[int] = Query(None, ge=1, le=3, description="Сложность (1-3)"),
    db: AsyncSession = Depends(get_db),
):
    """Агрегация количества билетов по категориям и подкатегориям."""
    stmt = select(
        GeneratedTicket.category, 
        GeneratedTicket.subcategory, 
        func.count(GeneratedTicket.id).label("count")
    )
    if complexity is not None:
        stmt = stmt.where(GeneratedTicket.complexity == complexity)

    stmt = stmt.group_by(GeneratedTicket.category, GeneratedTicket.subcategory)
    
    result = await db.execute(stmt)
    records = result.all()
    
    counts = {}
    for cat, subcat, count in records:
        cat_name = cat or "Общее"
        if cat_name not in counts:
            counts[cat_name] = {"category": cat_name, "total": 0, "subcategories": {}}
        counts[cat_name]["total"] += count
        
        if subcat:
            counts[cat_name]["subcategories"][subcat] = count
            
    return list(counts.values())

@api_tickets_router.get("/{ticket_id}", response_model=TicketResponse)
@api_tickets_router.get("/{ticket_id}/", response_model=TicketResponse, include_in_schema=False)
@tickets_router.get("/{ticket_id}", response_model=TicketResponse)
@tickets_router.get("/{ticket_id}/", response_model=TicketResponse, include_in_schema=False)
async def get_ticket(
    ticket_id: str,
    db: AsyncSession = Depends(get_db),
) -> TicketResponse:
    # Try GeneratedTicket
    ticket = await db.get(GeneratedTicket, ticket_id)
    if ticket:
        seq = ticket.sequence_number if ticket.sequence_number is not None else 0
        display_id = f"{make_abbr(ticket.category)}_{make_abbr(ticket.subcategory)}_{seq}"
        return TicketResponse(
            id=ticket.id,
            ticket_id=ticket.id,
            display_id=display_id,
            sequence_number=ticket.sequence_number,
            category=ticket.category,
            subcategory=ticket.subcategory,
            complexity=ticket.complexity,
            plot=ticket.plot,
            factoids=ticket.factoids or {},
            ground_truth=ticket.ground_truth or {},
            etalon_services=ticket.etalon_services or [],
            status="active",
            created_at=ticket.created_at,
        )
    
    # Try ScenarioTicket
    scenario = await db.get(ScenarioTicket, ticket_id)
    if scenario:
        settings = scenario.settings or {}
        ai = scenario.ai_content or {}
        gt = scenario.ground_truth or {}
        cat = settings.get("category", "Общее")
        sub = settings.get("subcategory")
        seq = settings.get("sequence_number", 0)
        display_id = f"{make_abbr(cat)}_{make_abbr(sub)}_{seq}"
        return TicketResponse(
            id=scenario.scenario_id,
            ticket_id=scenario.scenario_id,
            display_id=display_id,
            sequence_number=seq,
            category=cat,
            subcategory=sub,
            complexity=settings.get("complexity", 1),
            plot=settings.get("plot") or ai.get("plot", ""),
            factoids=ai.get("factoids", {}),
            ground_truth=gt,
            etalon_services=settings.get("etalon_services", []),
            status=scenario.workflow_state.get("status", "active") if scenario.workflow_state else "active",
            created_at=scenario.created_at,
        )
        
    raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Билет не найден")



@api_tickets_router.get("", response_model=List[TicketResponse])
@api_tickets_router.get("/", response_model=List[TicketResponse], include_in_schema=False)
@tickets_router.get("", response_model=List[TicketResponse])
@tickets_router.get("/", response_model=List[TicketResponse], include_in_schema=False)
async def list_tickets(
    category: Optional[str] = Query(None, description="Категория билета"),
    subcategory: Optional[str] = Query(None, description="Подкатегория"),
    complexity: Optional[int] = Query(None, ge=1, le=3, description="Сложность (1-3)"),
    limit: Optional[int] = Query(100, ge=1, le=500),
    include_deleted: Optional[bool] = Query(False, description="Включая удаленные"),
    db: AsyncSession = Depends(get_db),
) -> List[TicketResponse]:
    """Получение списка сохраненных билетов с поддержкой фильтрации."""
    TicketFilterParams(category=category, subcategory=subcategory, complexity=complexity)

    stmt = select(GeneratedTicket).order_by(desc(GeneratedTicket.created_at))
    if category:
        stmt = stmt.where(GeneratedTicket.category == category)
    if subcategory:
        stmt = stmt.where(GeneratedTicket.subcategory == subcategory)
    if complexity is not None:
        stmt = stmt.where(GeneratedTicket.complexity == complexity)
    
    if not include_deleted:
        from sqlalchemy import or_
        stmt = stmt.where(or_(GeneratedTicket.is_deleted == False, GeneratedTicket.is_deleted.is_(None)))

    if limit is not None:
        stmt = stmt.limit(limit)

    result = await db.execute(stmt)
    records = result.scalars().all()

    output: List[TicketResponse] = []
    for r in records:
        seq = r.sequence_number if r.sequence_number is not None else 0
        display_id = f"{make_abbr(r.category)}_{make_abbr(r.subcategory)}_{seq}"
        output.append(
            TicketResponse(
                id=r.id,
                ticket_id=r.id,
                display_id=display_id,
                sequence_number=r.sequence_number,
                category=r.category,
                subcategory=r.subcategory,
                complexity=r.complexity,
                plot=r.plot,
                factoids=r.factoids or {},
                ground_truth=r.ground_truth or {},
                etalon_services=r.etalon_services or [],
                status="active",
                created_at=r.created_at,
            )
        )

    # Legacy fallback to ScenarioTicket if GeneratedTicket has no records and no filters
    if not output and category is None and subcategory is None and complexity is None:
        stmt_legacy = select(ScenarioTicket).order_by(desc(ScenarioTicket.created_at))
        
        if not include_deleted:
            from sqlalchemy import or_
            stmt_legacy = stmt_legacy.where(or_(ScenarioTicket.is_deleted == False, ScenarioTicket.is_deleted.is_(None)))
            
        stmt_legacy = stmt_legacy.limit(limit or 100)
        res_legacy = await db.execute(stmt_legacy)
        for r in res_legacy.scalars().all():
            settings = r.settings or {}
            ai = r.ai_content or {}
            gt = r.ground_truth or {}
            cat = settings.get("category", "Общее")
            sub = settings.get("subcategory")
            seq = settings.get("sequence_number", 0)
            display_id = f"{make_abbr(cat)}_{make_abbr(sub)}_{seq}"
            output.append(
                TicketResponse(
                    id=r.scenario_id,
                    ticket_id=r.scenario_id,
                    display_id=display_id,
                    sequence_number=seq,
                    category=cat,
                    subcategory=sub,
                    complexity=settings.get("complexity", 1),
                    plot=settings.get("plot") or ai.get("plot", ""),
                    factoids=ai.get("factoids", {}),
                    ground_truth=gt,
                    etalon_services=settings.get("etalon_services", []),
                    status=r.workflow_state.get("status", "active") if r.workflow_state else "active",
                    created_at=r.created_at,
                )
            )

    return output


@api_tickets_router.post(
    "/generate",
    status_code=status.HTTP_202_ACCEPTED,
    response_model=TicketGenerateAcceptedResponse,
)
@api_tickets_router.post(
    "/generate/",
    status_code=status.HTTP_202_ACCEPTED,
    response_model=TicketGenerateAcceptedResponse,
    include_in_schema=False,
)
@tickets_router.post(
    "/generate",
    status_code=status.HTTP_202_ACCEPTED,
    response_model=TicketGenerateAcceptedResponse,
)
@tickets_router.post(
    "/generate/",
    status_code=status.HTTP_202_ACCEPTED,
    response_model=TicketGenerateAcceptedResponse,
    include_in_schema=False,
)
async def generate_ticket_endpoint(
    req: TicketGenerateRequest,
    background_tasks: BackgroundTasks,
) -> TicketGenerateAcceptedResponse:
    """Асинхронная генерация пакета билетов через BackgroundTasks."""
    global active_generations
    active_generations += req.count
    background_tasks.add_task(
        generate_tickets_background_task,
        category=req.category,
        subcategory=req.subcategory,
        count=req.count,
    )
    return TicketGenerateAcceptedResponse(message="Генерация начата")


def generate_tickets_import_template_excel() -> bytes:
    """Генерация Excel-шаблона для импорта билетов."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Шаблон билетов"

    headers = [
        "Название",
        "Описание",
        "Номер звонящего",
        "Текст абонента",
        "Целевая служба (01, 02)",
        "Обязательные фактоиды (через запятую)",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # Демонстрационные примеры строк
    sample_rows = [
        [
            "Пожар в жилом доме",
            "Возгорание на кухне на 3 этаже, сильный дым",
            "+79161234567",
            "Здравствуйте! Горит квартира на 3 этаже, сильный дым в подъезде по ул. Ленина, д. 5! Срочно пожарных!",
            "01",
            "адрес, этаж, задымление, пострадавшие",
        ],
        [
            "ДТП на перекрестке",
            "Лобовое столкновение двух автомобилей, есть пострадавшие",
            "+79257654321",
            "Авария на перекрестке ул. Мира и Советской! Машины разбиты, водителю плохо!",
            "01, 02, 03",
            "адрес, автомобили, пострадавшие, скорая помощь",
        ],
    ]
    for row in sample_rows:
        ws.append(row)

    col_widths = {1: 28, 2: 36, 3: 20, 4: 55, 5: 26, 6: 40}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 26

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


@api_tickets_router.get("/import/template")
@api_tickets_router.get("/import/template/", include_in_schema=False)
@tickets_router.get("/import/template")
@tickets_router.get("/import/template/", include_in_schema=False)
async def download_tickets_import_template() -> Response:
    """Эндпоинт для скачивания Excel-шаблона импорта билетов."""
    file_bytes = generate_tickets_import_template_excel()
    return Response(
        content=file_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="tickets_template.xlsx"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


@api_tickets_router.post("/import/upload")
@api_tickets_router.post("/import/upload/", include_in_schema=False)
@tickets_router.post("/import/upload")
@tickets_router.post("/import/upload/", include_in_schema=False)
async def upload_tickets_import(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
) -> Dict[str, Any]:
    """Эндпоинт для загрузки и пакетного импорта билетов из Excel."""
    filename = (file.filename or "").lower()
    if not (filename.endswith(".xlsx") or filename.endswith(".xls")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Неподдерживаемый формат файла. Требуется файл .xlsx или .xls",
        )

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Файл пуст",
        )

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Не удалось распарсить Excel файл: {str(e)}",
        )

    ws = wb.active
    if ws is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="В файле не найден активный лист",
        )

    # Определяем строку заголовков и маппинг колонок
    header_row_idx = 1
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), start=1):
        row_strs = [str(c or "").strip().lower() for c in row if c is not None]
        if any("целев" in s or "служб" in s or "назван" in s or "текст" in s for s in row_strs):
            header_row_idx = r_idx
            break

    header_cells = list(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True))[0]
    col_map: Dict[str, int] = {}
    for idx, cell_val in enumerate(header_cells):
        val_str = str(cell_val or "").strip().lower()
        if "назван" in val_str or "заголовок" in val_str or "тема" in val_str:
            col_map["title"] = idx
        elif "описан" in val_str:
            col_map["desc"] = idx
        elif "номер" in val_str or "телефон" in val_str or "звонящ" in val_str:
            col_map["phone"] = idx
        elif "текст" in val_str or "абонент" in val_str or "сообщен" in val_str:
            col_map["caller_text"] = idx
        elif "целев" in val_str or "служб" in val_str:
            col_map["services"] = idx
        elif "фактоид" in val_str:
            col_map["factoids"] = idx

    # Позиционный фоллбэк
    if "title" not in col_map and len(header_cells) > 0:
        col_map["title"] = 0
    if "desc" not in col_map and len(header_cells) > 1:
        col_map["desc"] = 1
    if "phone" not in col_map and len(header_cells) > 2:
        col_map["phone"] = 2
    if "caller_text" not in col_map and len(header_cells) > 3:
        col_map["caller_text"] = 3
    if "services" not in col_map and len(header_cells) > 4:
        col_map["services"] = 4
    if "factoids" not in col_map and len(header_cells) > 5:
        col_map["factoids"] = 5

    def get_cell_str(row_data: tuple, key: str) -> str:
        idx = col_map.get(key)
        if idx is not None and idx < len(row_data) and row_data[idx] is not None:
            return str(row_data[idx]).strip()
        return ""

    errors: List[str] = []
    success_count = 0

    # Ensure sys_sequences table exists
    await db.execute(
        text("CREATE TABLE IF NOT EXISTS sys_sequences (name VARCHAR PRIMARY KEY, last_val INTEGER NOT NULL DEFAULT 0)")
    )
    try:
        await db.execute(
            text("INSERT INTO sys_sequences (name, last_val) VALUES ('generated_tickets', 0) ON CONFLICT (name) DO NOTHING")
        )
    except Exception:
        await db.execute(
            text("INSERT OR IGNORE INTO sys_sequences (name, last_val) VALUES ('generated_tickets', 0)")
        )
    seq_res = await db.execute(text("SELECT last_val FROM sys_sequences WHERE name='generated_tickets'"))
    current_seq = seq_res.scalar() or 0

    for row_idx, row in enumerate(
        ws.iter_rows(min_row=header_row_idx + 1, values_only=True),
        start=header_row_idx + 1,
    ):
        title = get_cell_str(row, "title")
        desc_text = get_cell_str(row, "desc")
        phone = get_cell_str(row, "phone")
        caller_text = get_cell_str(row, "caller_text")
        services_str = get_cell_str(row, "services")
        factoids_str = get_cell_str(row, "factoids")

        # Пропускаем полностью пустые строки
        if not any([title, desc_text, phone, caller_text, services_str, factoids_str]):
            continue

        # Валидация: целевая служба обязательна
        if not services_str:
            errors.append(f"строка {row_idx} пустая целевая служба")
            continue

        # Валидация: хотя бы название или текст абонента
        if not title and not caller_text:
            errors.append(f"строка {row_idx} пустое название или текст абонента")
            continue

        if not title:
            title = (caller_text[:47] + "...") if len(caller_text) > 50 else caller_text
        if not caller_text:
            caller_text = desc_text or title
        if not desc_text:
            desc_text = caller_text

        # Извлечение служб
        extracted_services: List[str] = []
        lower_s = services_str.lower()
        has_01 = bool(re.search(r"\b(01|101)\b", lower_s) or "пожар" in lower_s)
        has_02 = bool(re.search(r"\b(02|102)\b", lower_s) or "полиц" in lower_s or "мвд" in lower_s or "дпс" in lower_s)
        has_03 = bool(re.search(r"\b(03|103)\b", lower_s) or "скор" in lower_s or "медиц" in lower_s or "врач" in lower_s)
        has_04 = bool(re.search(r"\b(04|104)\b", lower_s) or "газ" in lower_s)

        if has_01:
            extracted_services.append("01 Пожарные")
        if has_02:
            extracted_services.append("02 Полиция")
        if has_03:
            extracted_services.append("03 Скорая")
        if has_04:
            extracted_services.append("04 Газ")

        if not extracted_services:
            for part in re.split(r"[,;]+", services_str):
                p = part.strip()
                if p:
                    extracted_services.append(p)

        # Разбор фактоидов
        factoids_list = [f.strip() for f in re.split(r"[,;\n]+", factoids_str) if f.strip()]
        factoids_dict: Dict[str, Any] = {f"factoid_{i+1}": f for i, f in enumerate(factoids_list)}
        if caller_text:
            factoids_dict["situation_1"] = caller_text

        # Определение категории
        combined_text = f"{title} {desc_text} {caller_text}".lower()
        if "01 Пожарные" in extracted_services or "пожар" in combined_text or "дым" in combined_text:
            category = "Пожары и задымления"
        elif "04 Газ" in extracted_services or "газ" in combined_text:
            category = "Запах газа"
        elif "02 Полиция" in extracted_services or "дтп" in combined_text or "авари" in combined_text or "столкнов" in combined_text:
            category = "ДТП" if ("дтп" in combined_text or "авари" in combined_text or "столкнов" in combined_text) else "Нарушение правопорядка"
        elif "03 Скорая" in extracted_services or "медиц" in combined_text or "сердц" in combined_text or "ранен" in combined_text:
            category = "Оказание медицинской скорой и неотложной помощи"
        else:
            category = "Общее"

        current_seq += 1
        ticket_id = f"import-{uuid.uuid4().hex[:12]}"
        plot = caller_text
        ground_truth: Dict[str, Any] = {"phone": phone} if phone else {}

        # Создаем GeneratedTicket (для таблицы билетов)
        gen_ticket = GeneratedTicket(
            id=ticket_id,
            category=category,
            subcategory=title,
            complexity=1,
            plot=plot,
            factoids=factoids_dict,
            ground_truth=ground_truth,
            etalon_services=extracted_services,
            sequence_number=current_seq,
            status="active",
        )
        db.add(gen_ticket)

        # Создаем ScenarioTicket (требование ТЗ)
        scenario_ticket = ScenarioTicket(
            scenario_id=ticket_id,
            title=title,
            category=category,
            complexity=1,
            content={
                "title": title,
                "description": desc_text,
                "caller_text": caller_text,
            },
            settings={
                "complexity": 1,
                "etalon_services": extracted_services,
                "plot": plot,
                "category": category,
                "subcategory": title,
                "sequence_number": current_seq,
            },
            ground_truth=ground_truth,
            ai_content={
                "factoids": factoids_dict,
                "plot": plot,
            },
            workflow_state={"status": "active"},
        )
        db.add(scenario_ticket)
        success_count += 1

    if success_count > 0:
        await db.execute(
            text("UPDATE sys_sequences SET last_val = :val WHERE name = 'generated_tickets'"),
            {"val": current_seq},
        )
        await db.commit()

    return {
        "success": len(errors) == 0 or success_count > 0,
        "success_count": success_count,
        "error_count": len(errors),
        "errors": errors,
        "total_rows": success_count + len(errors),
    }


def _generate_preview_audio_for_ticket(plot: str, gt: dict):
    from backend.core.text_normalization import expand_address_for_tts, format_phone_for_tts
    from backend.core.tts_v2 import tts_engine_v2
    if not isinstance(gt, dict):
        gt = {}
    speaker = gt.get("speaker", "aidar")
    texts = []
    if plot: texts.append(str(plot))
    if gt.get("fio"): texts.append(str(gt["fio"]))
    if gt.get("phone"): texts.append("номер " + format_phone_for_tts(str(gt["phone"])))
    address_parts = []
    def clean_val(k):
        v = gt.get(k)
        if v is None or str(v).lower().strip() in ['none', 'null', '0', '-', '']: return ""
        return str(v).strip()
    street = clean_val("street")
    if street:
        if not any(m in street.lower() for m in ['ул.', 'улица', 'ш.', 'шоссе', 'пр-кт', 'проспект', 'пер.', 'переулок', 'бульвар', 'б-р']):
            address_parts.append(f"ул. {street}")
        else: address_parts.append(street)
    house = clean_val("house")
    if house:
        if not any(m in house.lower() for m in ['д.', 'дом', 'стр', 'строение', 'корп', 'корпус', 'влад']):
            address_parts.append(f"д. {house}")
        else: address_parts.append(house)
    entrance = clean_val("entrance")
    if entrance: address_parts.append(f"под. {entrance}")
    floor = clean_val("floor")
    if floor: address_parts.append(f"эт. {floor}")
    apartment = clean_val("apartment")
    if apartment:
        if not any(m in apartment.lower() for m in ['кв', 'квартира', 'комн']): address_parts.append(f"кв. {apartment}")
        else: address_parts.append(apartment)
    intercom = clean_val("intercom")
    if intercom: address_parts.append(f"домофон {intercom}")
    if address_parts:
        texts.append(expand_address_for_tts(", ".join(address_parts)))
    return tts_engine_v2.concatenate_tts(texts, speaker=speaker)

@api_tickets_router.get("/{ticket_id}/audio")
@api_tickets_router.get("/{ticket_id}/audio/", include_in_schema=False)
@tickets_router.get("/{ticket_id}/audio")
@tickets_router.get("/{ticket_id}/audio/", include_in_schema=False)
async def get_ticket_audio(
    ticket_id: str,
    db: AsyncSession = Depends(get_db),
) -> Response:
    """Генерация/склеивание аудио озвучки билета."""
    try:
        ticket = await db.get(GeneratedTicket, ticket_id)
        plot = ""
        gt = {}
        if ticket:
            plot = ticket.plot or ""
            gt = ticket.ground_truth or {}
        else:
            scenario_ticket = await db.get(ScenarioTicket, ticket_id)
            if not scenario_ticket:
                raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Билет не найден")
            settings = scenario_ticket.settings or {}
            ai = scenario_ticket.ai_content or {}
            plot = settings.get("plot") or ai.get("plot", "")
            gt = scenario_ticket.ground_truth or {}

        audio_bytes = await asyncio.to_thread(_generate_preview_audio_for_ticket, plot, gt)
        return Response(content=audio_bytes, media_type="audio/wav")
    except HTTPException:
        raise
    except asyncio.CancelledError:
        logger.warning("Ticket audio generation cancelled by client for ticket %s", ticket_id)
        raise
    except Exception as e:
        logger.error("Error generating ticket audio for ticket %s: %s", ticket_id, e, exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@api_tickets_router.get("/tts/speak")
@tickets_router.get("/tts/speak")
async def tts_speak(
    text: str = Query(..., description="Текст для озвучки"),
    speaker: str = Query("aidar", description="Голос TTS (aidar, baya, kseniya, eugene)"),
) -> Response:
    """Генерирует WAV-аудио для произвольного текста через TTS-движок.
    Используется фронтендом для воспроизведения реплик заявителя по одной."""
    try:
        if not text or not text.strip():
            raise HTTPException(status_code=400, detail="Параметр 'text' не может быть пустым")
        audio_bytes = await asyncio.to_thread(tts_engine_v2.concatenate_tts, [text.strip()], speaker)
        return Response(
            content=audio_bytes,
            media_type="audio/wav",
            headers={"Cache-Control": "no-cache"},
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error("Error in tts_speak: %s", e, exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@api_tickets_router.get("/tts/service-reply")
@tickets_router.get("/tts/service-reply")
async def get_service_reply(
    service_name: str,
):
    try:
        service_lower = service_name.lower()
        if "пожар" in service_lower or "101" in service_lower or "мчс" in service_lower:
            phrase = "Пожарная охрана, ... радиотелефонист Иванова Мария Сергеевна, ... слушаю."
            speaker = "baya"
        elif "полиц" in service_lower or "102" in service_lower or "гибдд" in service_lower:
            phrase = "Дежурная часть полиции, ... майор Смирнов Петр Алексеевич, ... слушаю."
            speaker = "aidar"
        elif "скор" in service_lower or "103" in service_lower or "мц" in service_lower:
            phrase = "Станция скорой медицинской помощи, ... диспетчер Соколова Анна Юрьевна, ... слушаю."
            speaker = "kseniya"
        elif "газ" in service_lower or "104" in service_lower:
            phrase = "Аварийная служба газа, ... мастер Петров Иван Васильевич, ... слушаю."
            speaker = "eugene"
        else:
            phrase = f"Дежурный диспетчер службы {service_name}, ... слушаю."
            speaker = "baya"
            
        audio_bytes = await asyncio.to_thread(tts_engine_v2.concatenate_tts, [phrase], speaker)
        return Response(content=audio_bytes, media_type="audio/wav")
    except Exception as e:
        logger.error("Error generating service reply audio: %s", e)
        raise HTTPException(status_code=500, detail=str(e))

@api_tickets_router.get("/tts/service-accepted")
@tickets_router.get("/tts/service-accepted")
async def get_service_accepted(
    service_name: str,
):
    try:
        service_lower = service_name.lower()
        if "пожар" in service_lower or "101" in service_lower or "мчс" in service_lower:
            speaker = "baya"
            phrase = "Информацию, ... приняла."
        elif "полиц" in service_lower or "102" in service_lower or "гибдд" in service_lower:
            speaker = "aidar"
            phrase = "Информацию, ... принял."
        elif "скор" in service_lower or "103" in service_lower or "мц" in service_lower:
            speaker = "kseniya"
            phrase = "Информацию, ... приняла."
        elif "газ" in service_lower or "104" in service_lower:
            speaker = "eugene"
            phrase = "Информацию, ... принял."
        else:
            speaker = "baya"
            phrase = "Информацию, ... приняла."
        audio_bytes = await asyncio.to_thread(tts_engine_v2.concatenate_tts, [phrase], speaker)
        return Response(content=audio_bytes, media_type="audio/wav")
    except Exception as e:
        logger.error("Error generating service accepted audio: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


@api_tickets_router.patch("/{ticket_id}", response_model=TicketResponse)
@api_tickets_router.patch("/{ticket_id}/", response_model=TicketResponse, include_in_schema=False)
@tickets_router.patch("/{ticket_id}", response_model=TicketResponse)
@tickets_router.patch("/{ticket_id}/", response_model=TicketResponse, include_in_schema=False)
async def update_ticket(
    ticket_id: str,
    ticket_in: TicketUpdate,
    db: AsyncSession = Depends(get_db),
) -> TicketResponse:
    ticket = await db.get(GeneratedTicket, ticket_id)
    scenario_ticket = await db.get(ScenarioTicket, ticket_id)
    if not ticket and not scenario_ticket:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Билет не найден")
    
    update_data = ticket_in.model_dump(exclude_unset=True)
    if ticket:
        for field, value in update_data.items():
            setattr(ticket, field, value)
        await db.commit()
        await db.refresh(ticket)
        seq = ticket.sequence_number or 0
        display_id = f"{make_abbr(ticket.category)}_{make_abbr(ticket.subcategory)}_{seq}"
        res_cat = ticket.category
        res_sub = ticket.subcategory
        res_comp = ticket.complexity
        res_plot = ticket.plot
        res_facts = ticket.factoids
        res_gt = ticket.ground_truth
        res_serv = ticket.etalon_services
        res_created = ticket.created_at
    else:
        # Fallback if only scenario_ticket exists
        seq = (scenario_ticket.settings or {}).get("sequence_number", 0)
        res_cat = update_data.get("category", (scenario_ticket.settings or {}).get("category", "Общее"))
        res_sub = update_data.get("subcategory", (scenario_ticket.settings or {}).get("subcategory"))
        display_id = f"{make_abbr(res_cat)}_{make_abbr(res_sub)}_{seq}"
        res_comp = update_data.get("complexity", (scenario_ticket.settings or {}).get("complexity", 1))
        res_plot = update_data.get("plot", (scenario_ticket.settings or {}).get("plot", ""))
        res_facts = update_data.get("factoids", (scenario_ticket.ai_content or {}).get("factoids", {}))
        res_gt = update_data.get("ground_truth", scenario_ticket.ground_truth or {})
        res_serv = update_data.get("etalon_services", (scenario_ticket.settings or {}).get("etalon_services", []))
        res_created = scenario_ticket.created_at

    if scenario_ticket:
        settings = dict(scenario_ticket.settings or {})
        for field in ["category", "subcategory", "complexity", "plot", "etalon_services"]:
            if field in update_data:
                settings[field] = update_data[field]
        scenario_ticket.settings = settings
        if "ground_truth" in update_data:
            scenario_ticket.ground_truth = update_data["ground_truth"]
        if "plot" in update_data:
            ai_content = dict(scenario_ticket.ai_content or {})
            ai_content["plot"] = update_data["plot"]
            scenario_ticket.ai_content = ai_content
        await db.commit()
    
    return TicketResponse(
        id=ticket_id,
        ticket_id=ticket_id,
        display_id=display_id,
        sequence_number=seq,
        category=res_cat,
        subcategory=res_sub,
        complexity=res_comp,
        plot=res_plot,
        factoids=res_facts,
        ground_truth=res_gt,
        etalon_services=res_serv,
        status="active",
        created_at=res_created,
    )

@api_tickets_router.delete("/{ticket_id}", status_code=status.HTTP_200_OK)
@api_tickets_router.delete("/{ticket_id}/", status_code=status.HTTP_200_OK, include_in_schema=False)
@tickets_router.delete("/{ticket_id}", status_code=status.HTTP_200_OK)
@tickets_router.delete("/{ticket_id}/", status_code=status.HTTP_200_OK, include_in_schema=False)
async def delete_ticket(
    ticket_id: str,
    db: AsyncSession = Depends(get_db),
) -> Dict[str, Any]:
    """Удаление билета по ID из GeneratedTicket и ScenarioTicket."""
    gen_ticket = await db.get(GeneratedTicket, ticket_id)
    scenario_ticket = await db.get(ScenarioTicket, ticket_id)

    if not gen_ticket and not scenario_ticket:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Билет не найден",
        )

    if gen_ticket:
        gen_ticket.is_deleted = True
    if scenario_ticket:
        scenario_ticket.is_deleted = True

    await db.commit()
    return {"status": "ok", "message": f"Билет {ticket_id} успешно удален", "id": ticket_id}


class TicketRefineRequest(BaseModel):
    correction_comment: str

@api_tickets_router.post("/{ticket_id}/refine", response_model=TicketResponse)
@api_tickets_router.post("/{ticket_id}/refine/", response_model=TicketResponse, include_in_schema=False)
@tickets_router.post("/{ticket_id}/refine", response_model=TicketResponse)
@tickets_router.post("/{ticket_id}/refine/", response_model=TicketResponse, include_in_schema=False)
async def refine_ticket(
    ticket_id: str,
    req: TicketRefineRequest,
    db: AsyncSession = Depends(get_db),
) -> TicketResponse:
    """Перегенерация сюжета билета на основе комментария (заглушка)."""
    gen_ticket = await db.get(GeneratedTicket, ticket_id)
    scenario_ticket = await db.get(ScenarioTicket, ticket_id)

    if not gen_ticket and not scenario_ticket:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Билет не найден")
    
    # Имитация вызова LLM
    await asyncio.sleep(2)
    
    plot_addition = f"\n[ИИ-коррекция с учетом: {req.correction_comment}]"
    
    if gen_ticket:
        gen_ticket.plot = (gen_ticket.plot or "") + plot_addition
        seq = gen_ticket.sequence_number or 0
        display_id = f"{make_abbr(gen_ticket.category)}_{make_abbr(gen_ticket.subcategory)}_{seq}"
        res_cat = gen_ticket.category
        res_sub = gen_ticket.subcategory
        res_comp = gen_ticket.complexity
        res_plot = gen_ticket.plot
        res_facts = gen_ticket.factoids or {}
        res_gt = gen_ticket.ground_truth or {}
        res_serv = gen_ticket.etalon_services or []
        res_created = gen_ticket.created_at

    if scenario_ticket:
        settings = dict(scenario_ticket.settings or {})
        ai_content = dict(scenario_ticket.ai_content or {})
        
        current_plot = settings.get("plot") or ai_content.get("plot", "")
        new_plot = current_plot + plot_addition
        
        settings["plot"] = new_plot
        ai_content["plot"] = new_plot
        
        scenario_ticket.settings = settings
        scenario_ticket.ai_content = ai_content

        if not gen_ticket:
            seq = settings.get("sequence_number", 0)
            res_cat = settings.get("category", "Общее")
            res_sub = settings.get("subcategory")
            display_id = f"{make_abbr(res_cat)}_{make_abbr(res_sub)}_{seq}"
            res_comp = settings.get("complexity", 1)
            res_plot = new_plot
            res_facts = ai_content.get("factoids", {})
            res_gt = scenario_ticket.ground_truth or {}
            res_serv = settings.get("etalon_services", [])
            res_created = scenario_ticket.created_at

    await db.commit()
    
    return TicketResponse(
        id=ticket_id,
        ticket_id=ticket_id,
        display_id=display_id,
        sequence_number=seq,
        category=res_cat,
        subcategory=res_sub,
        complexity=res_comp,
        plot=res_plot,
        factoids=res_facts,
        ground_truth=res_gt,
        etalon_services=res_serv,
        status="active",
        created_at=res_created,
    )


@api_tickets_router.get("/addresses")
@api_tickets_router.get("/addresses/", include_in_schema=False)
@tickets_router.get("/addresses")
@tickets_router.get("/addresses/", include_in_schema=False)
async def search_addresses(q: str = Query("", description="Поисковый запрос")):
    """Поиск адресов по строке."""
    project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    possible_paths = [
        os.path.join(project_root, "data", "moscow_112_addresses.json"),
        os.path.join(os.getcwd(), "data", "moscow_112_addresses.json"),
    ]
    data = []
    for path in possible_paths:
        if os.path.exists(path):
            try:
                with open(path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                break
            except Exception as e:
                logger.warning("Error loading moscow_112_addresses.json from %s: %s", path, e)
    
    q_lower = q.lower()
    results = []
    seen = set()
    for item in data:
        street = item.get("street", "")
        if q_lower in street.lower():
            okrug = item.get("okrug", "")
            district = item.get("district", "")
            key = f"{street}|{okrug}|{district}"
            if key not in seen:
                seen.add(key)
                results.append({"street": street, "okrug": okrug, "district": district})
                if len(results) >= 20:
                    break
    
    return results
