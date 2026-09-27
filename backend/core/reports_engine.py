"""Core logic for calculating evaluation metrics, scores, and generating reports."""

import csv
import io
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Tuple
from jinja2 import Environment, FileSystemLoader, select_autoescape


def calculate_final_score(
    comm_metrics: Dict[str, Any],
    card_metrics: Dict[str, Any],
    sla_metrics: Dict[str, Any],
) -> Tuple[int, Dict[str, Any]]:
    """Calculate the final evaluation score and penalty details from 3 blocks.

    Blocks:
    1. Communication:
       - greeting_success: penalty 10 if False ('no_greeting')
       - filler_words: penalty min(20, count * 2) if > 0 ('filler_words')
       - script_followed_pct: penalty based on deviation, 'script_violation' if < 80%
    2. Incident Card:
       - address_correct: penalty 15 if False ('address_error')
       - services_matched: penalty 10 if False ('services_mismatched')
       - over_dispatched_services: penalty 15 per extra service ('over_dispatching')
       - missed_critical_factoids: penalty 20 per missed factoid ('missed_factoid')
    3. SLA:
       - sla_breached_count: penalty 10 per breach ('sla_breach')
       - time_to_first_dispatch_sec: penalty 15 if > 60s ('slow_dispatch')

    Returns:
        (final_score: int, details: dict)
    """
    penalties_list: List[str] = []
    breakdown: Dict[str, int] = {}

    # ─── 1. Communication block ───────────────────────────────────────────────
    comm_penalty = 0
    greeting_success = comm_metrics.get("greeting_success", True)
    if not greeting_success:
        comm_penalty += 10
        penalties_list.append("no_greeting")
        breakdown["no_greeting"] = 10

    filler_words = comm_metrics.get("filler_words", 0)
    if filler_words > 0:
        fw_pen = min(20, int(filler_words * 2))
        comm_penalty += fw_pen
        penalties_list.append("filler_words")
        breakdown["filler_words"] = fw_pen

    script_followed = comm_metrics.get("script_followed_pct", 100)
    if script_followed < 100:
        sf_pen = int(round((100 - script_followed) * 0.3))
        comm_penalty += sf_pen
        breakdown["script_violation"] = sf_pen
        if script_followed < 80:
            penalties_list.append("script_violation")

    # ─── 2. Incident Card block ───────────────────────────────────────────────
    card_penalty = 0
    address_correct = card_metrics.get("address_correct", True)
    if not address_correct:
        card_penalty += 15
        penalties_list.append("address_error")
        breakdown["address_error"] = 15

    services_matched = card_metrics.get("services_matched", True)
    if not services_matched:
        card_penalty += 10
        penalties_list.append("services_mismatched")
        breakdown["services_mismatched"] = 10

    over_dispatched = card_metrics.get("over_dispatched_services") or []
    if over_dispatched:
        od_pen = 15 * len(over_dispatched)
        card_penalty += od_pen
        penalties_list.append("over_dispatching")
        breakdown["over_dispatching"] = od_pen

    missed_factoids = card_metrics.get("missed_critical_factoids") or []
    if missed_factoids:
        mf_pen = 20 * len(missed_factoids)
        card_penalty += mf_pen
        penalties_list.append("missed_factoid")
        breakdown["missed_factoid"] = mf_pen

    # ─── 3. SLA block ─────────────────────────────────────────────────────────
    sla_penalty = 0
    sla_breaches = sla_metrics.get("sla_breached_count", 0)
    if sla_breaches > 0:
        sb_pen = 10 * sla_breaches
        sla_penalty += sb_pen
        penalties_list.append("sla_breach")
        breakdown["sla_breach"] = sb_pen

    time_to_first = sla_metrics.get("time_to_first_dispatch_sec", 0)
    if time_to_first > 60:
        tt_pen = 15
        sla_penalty += tt_pen
        penalties_list.append("slow_dispatch")
        breakdown["slow_dispatch"] = tt_pen

    # ─── Aggregation ──────────────────────────────────────────────────────────
    total_penalty = comm_penalty + card_penalty + sla_penalty
    final_score = max(0, min(100, 100 - total_penalty))
    comm_score = max(0, min(100, 100 - comm_penalty))
    card_score = max(0, min(100, 100 - card_penalty))
    sla_score = max(0, min(100, 100 - sla_penalty))

    details = {
        "penalty": total_penalty,
        "penalties_list": penalties_list,
        "comm_score": comm_score,
        "card_score": card_score,
        "sla_score": sla_score,
        "final_score": final_score,
        "penalties": breakdown,
    }

    return final_score, details


def generate_pdf_from_html(html_content: str, output_path: str) -> str:
    """Convert HTML string to PDF file using weasyprint with xhtml2pdf fallback.

    Ensures the target directory exists and writes output to output_path.
    """
    out_file = Path(output_path).resolve()
    out_file.parent.mkdir(parents=True, exist_ok=True)

    try:
        import weasyprint
        weasyprint.HTML(string=html_content).write_pdf(str(out_file))
        return str(out_file)
    except Exception as wp_exc:
        # Fallback to xhtml2pdf if weasyprint encounters issues
        try:
            from xhtml2pdf import pisa
            with open(out_file, "wb") as f:
                res = pisa.CreatePDF(html_content, dest=f, encoding="utf-8")
                if res.err:
                    raise RuntimeError(f"xhtml2pdf error: {res.err}") from wp_exc
            return str(out_file)
        except Exception as xh_exc:
            raise RuntimeError(
                f"Failed to generate PDF. WeasyPrint: {wp_exc}; xhtml2pdf: {xh_exc}"
            ) from xh_exc


def get_templates_dir() -> Path:
    """Resolve backend/templates directory."""
    current_dir = Path(__file__).resolve().parent
    templates_dir = current_dir.parent / "templates"
    if not templates_dir.exists():
        templates_dir.mkdir(parents=True, exist_ok=True)
    return templates_dir


def render_report_html(context: Dict[str, Any], template_name: str = "report_layout.html") -> str:
    """Render the report HTML using Jinja2."""
    templates_dir = get_templates_dir()
    env = Environment(
        loader=FileSystemLoader(str(templates_dir)),
        autoescape=select_autoescape(["html", "xml"]),
    )
    template = env.get_template(template_name)
    return template.render(**context)


def generate_csv_from_data(cadets: List[Dict[str, Any]]) -> str:
    """Generate CSV string from cadet reports with UTF-8 BOM."""
    output = io.StringIO()
    # UTF-8 BOM for Excel compatibility
    output.write("\ufeff")
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)

    writer.writerow([
        "Курсант",
        "Билет",
        "Итоговый балл",
        "Балл коммуникация",
        "Балл карточка",
        "Балл SLA",
        "Время диспетчеризации (сек)",
        "Лишние службы (гипер-диспетчеризация)",
        "Пропущенные фактоиды",
        "Штрафы",
    ])

    for c in cadets:
        writer.writerow([
            c.get("cadet_name", "—"),
            c.get("ticket_id", "—"),
            c.get("final_score", 0),
            c.get("comm_score", 0),
            c.get("card_score", 0),
            c.get("sla_score", 0),
            c.get("time_to_first_dispatch_sec", 0),
            ", ".join(c.get("over_dispatched_services", [])) or "Нет",
            ", ".join(c.get("missed_critical_factoids", [])) or "Нет",
            ", ".join(c.get("penalties_list", [])) or "Нет",
        ])

    return output.getvalue()


def generate_excel_from_data(cadets_data: List[Dict[str, Any]]) -> bytes:
    """Generate a styled XLSX Excel spreadsheet from cadet reports in-memory.

    Args:
        cadets_data: List of cadet evaluation records.

    Returns:
        Raw bytes of the .xlsx file.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Аттестация курсантов"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    title_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")

    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    header_border = Border(
        left=Side(border_style="thin", color="1E40AF"),
        right=Side(border_style="thin", color="1E40AF"),
        top=Side(border_style="medium", color="1E3A8A"),
        bottom=Side(border_style="medium", color="1E3A8A"),
    )

    summary_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    summary_font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")

    # Row 1: Title Banner
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "СИСТЕМА-112: ОТЧЕТ ПО РЕЗУЛЬТАТАМ АТТЕСТАЦИИ И ТРЕНИРОВОК КУРСАНТОВ"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Row 2: Subtitle / Date
    ws.merge_cells("A2:J2")
    sub_cell = ws["A2"]
    now_str = datetime.now(timezone.utc).strftime("%d.%m.%Y %H:%M UTC")
    sub_cell.value = f"Сформировано: {now_str} | Всего записей: {len(cadets_data)}"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="64748B")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Row 3: Blank spacing row
    ws.row_dimensions[3].height = 10

    # Row 4: Column Headers
    headers = [
        "Курсант",
        "Билет / Задание",
        "Итоговый балл",
        "Балл коммуникация",
        "Балл карточка",
        "Балл SLA",
        "Время диспетчеризации (сек)",
        "Лишние службы (гипер-диспетчеризация)",
        "Пропущенные фактоиды",
        "Штрафы и замечания",
    ]

    ws.row_dimensions[4].height = 28
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    # Rows 5+: Cadet Data
    current_row = 5
    for idx, c in enumerate(cadets_data):
        row_fill = zebra_fill if idx % 2 == 1 else white_fill
        ws.row_dimensions[current_row].height = 22

        def _format_list(val: Any) -> str:
            if isinstance(val, (list, tuple, set)):
                return ", ".join(str(x) for x in val) if val else "—"
            return str(val) if val else "—"

        cadet_name = str(c.get("cadet_name") or "—")
        ticket_id = str(c.get("ticket_id") or "—")
        final_score = int(c.get("final_score", 0))
        comm_score = int(c.get("comm_score", 0))
        card_score = int(c.get("card_score", 0))
        sla_score = int(c.get("sla_score", 0))
        dispatch_sec = int(c.get("time_to_first_dispatch_sec", 0))
        over_disp = _format_list(c.get("over_dispatched_services"))
        missed_fact = _format_list(c.get("missed_critical_factoids"))
        penalties = _format_list(c.get("penalties_list"))

        row_values = [
            (cadet_name, Alignment(horizontal="left", vertical="center"), None),
            (ticket_id, Alignment(horizontal="center", vertical="center"), None),
            (final_score, Alignment(horizontal="center", vertical="center"), "score"),
            (comm_score, Alignment(horizontal="center", vertical="center"), None),
            (card_score, Alignment(horizontal="center", vertical="center"), None),
            (sla_score, Alignment(horizontal="center", vertical="center"), None),
            (dispatch_sec, Alignment(horizontal="center", vertical="center"), None),
            (over_disp, Alignment(horizontal="left", vertical="center", wrap_text=True), "alert" if over_disp != "—" else None),
            (missed_fact, Alignment(horizontal="left", vertical="center", wrap_text=True), "alert" if missed_fact != "—" else None),
            (penalties, Alignment(horizontal="left", vertical="center", wrap_text=True), None),
        ]

        for col_idx, (val, align, flag) in enumerate(row_values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.alignment = align
            cell.border = cell_border

            if flag == "score":
                if final_score >= 80:
                    cell.font = Font(name="Calibri", size=11, bold=True, color="15803D")
                    cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                elif final_score >= 60:
                    cell.font = Font(name="Calibri", size=11, bold=True, color="B45309")
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                else:
                    cell.font = Font(name="Calibri", size=11, bold=True, color="B91C1C")
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            elif flag == "alert":
                cell.font = Font(name="Calibri", size=10, color="B91C1C", bold=True)
                cell.fill = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
            else:
                cell.font = Font(name="Calibri", size=10, color="0F172A")
                cell.fill = row_fill

        current_row += 1

    # Summary Row if cadets_data has entries
    if cadets_data:
        ws.row_dimensions[current_row].height = 24
        avg_score = round(sum(c.get("final_score", 0) for c in cadets_data) / len(cadets_data), 1)
        ws.cell(row=current_row, column=1, value="Среднее значение:").font = summary_font
        ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=current_row, column=3, value=avg_score).font = summary_font
        ws.cell(row=current_row, column=3).alignment = Alignment(horizontal="center", vertical="center")

        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=c_idx)
            cell.fill = summary_fill
            cell.border = Border(
                top=Side(border_style="medium", color="2563EB"),
                bottom=Side(border_style="double", color="2563EB"),
                left=thin_border_side,
                right=thin_border_side,
            )
        current_row += 1

    # Freeze panes below headers
    ws.freeze_panes = "A5"

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in (1, 2):
                continue
            if cell.value:
                val_lines = str(cell.value).split("\n")
                line_max = max(len(l) for l in val_lines)
                if line_max > max_len:
                    max_len = line_max
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

