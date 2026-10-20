"""Read saved journal marks without conflating absent marks with zero or success."""
import io
import uuid
from datetime import datetime

import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient

from backend.api import router_reports
from backend.core.deps import get_current_user
from backend.core.gradebook_export import MOSCOW, build_gradebook_context, generate_gradebook_excel
from backend.core.reports_engine import render_report_html
from backend.main import app
from backend.models.domain_01 import User


RECORDS = [
    {"cadet_name": "Иванов Иван", "group_name": "Группа А", "lesson_title": "Занятие 1",
     "lesson_date": "2026-10-19T22:30:00Z", "ticket_id": "Лифт", "final_score": 0},
    {"cadet_name": "Петрова Анна", "group_name": "Группа А", "lesson_title": "Занятие 1",
     "lesson_date": "2026-10-20T10:00:00", "ticket_id": "Газ", "final_score": 94.5},
    {"cadet_name": "=Нет попытки", "group_name": "Группа Б", "lesson_title": "Занятие 2",
     "lesson_date": None, "ticket_id": "Нет попытки", "final_score": None},
]


def test_workbook_and_pdf_share_saved_marks_dates_and_missing_values():
    context = build_gradebook_context(RECORDS, datetime(2026, 10, 20, 12, 5, tzinfo=MOSCOW))
    assert context["avg_score"] == 47.2  # Includes zero, excludes the unmarked row.
    assert (context["graded_count"], context["pending_count"]) == (2, 1)
    workbook = openpyxl.load_workbook(io.BytesIO(generate_gradebook_excel(context)))
    sheet = workbook["Журнал оценок"]
    assert len(workbook.sheetnames) == 1
    assert "20.10.2026 12:05 МСК" in sheet["A2"].value
    assert "47.2" in sheet["A3"].value
    assert sheet["D5"].value == datetime(2026, 10, 20, 1, 30)
    assert sheet["D6"].value == datetime(2026, 10, 20, 13)
    assert sheet["D5"].number_format == "dd.mm.yyyy"
    assert sheet["D7"].value == "Не указана"
    assert [sheet.cell(r, 6).value for r in range(5, 8)] == [0, 94.5, None]
    assert [sheet.cell(r, 7).value for r in range(5, 8)] == ["Оценено", "Оценено", "Нет оценки"]
    assert sheet["A7"].data_type == "s"
    assert sheet["A7"].value == "'=Нет попытки"
    assert sheet.freeze_panes == "A5"
    assert sheet.auto_filter.ref == "A4:G7"
    assert sheet.page_setup.orientation == "landscape"
    html = render_report_html(context, "gradebook_layout.html")
    for text in ("Иванов Иван", "Петрова Анна", "94.5", "Нет оценки", "20.10.2026", "47.2"):
        assert text in html
    assert "Сдано" not in html
    assert "None" not in html


def test_ungraded_only_and_invalid_date_do_not_imply_zero_or_success():
    context = build_gradebook_context([{**RECORDS[2], "lesson_date": "not-a-date"}], datetime.now(MOSCOW))
    assert context["avg_score"] is None
    assert context["rows"][0]["lesson_date"] is None
    html = render_report_html(context, "gradebook_layout.html")
    assert "Нет оценок" in html
    assert "Нет оценки" in html


def test_legacy_pdf_accepts_unmarked_work():
    html = render_report_html({"generated_at": "20.10.2026", "cadets": [RECORDS[2]],
                               "critical_errors_count": 0, "over_dispatch_items": [], "missed_factoid_items": []})
    assert "Нет оценки" in html


@pytest.mark.asyncio
@pytest.mark.parametrize("format", ["excel", "pdf"])
async def test_generate_status_download_and_owner_boundary(format, monkeypatch, tmp_path):
    monkeypatch.setenv("REPORT_DIR", str(tmp_path))
    owner = User(user_id="teacher-gradebook", username="gradebook", role="TEACHER", is_active=True)
    app.dependency_overrides[get_current_user] = lambda: owner
    task_id = None
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.post("/api/reports/generate", json={
                "report_kind": "gradebook", "format": format, "custom_records": RECORDS,
            })
            assert response.status_code == 200, response.text
            task_id = response.json()["task_id"]
            status = await client.get(f"/api/reports/status/{task_id}")
            assert status.json()["status"] == "ready", status.text
            assert status.json()["generated_at"].endswith("+03:00")
            result = await client.get(f"/api/reports/download?task_id={task_id}&format={format}")
            assert result.status_code == 200, result.text
            if format == "pdf":
                assert result.headers["content-type"] == "application/pdf"
                assert result.content.startswith(b"%PDF-")
            else:
                assert result.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                workbook = openpyxl.load_workbook(io.BytesIO(result.content))
                assert workbook.active["F5"].value == 0
                assert workbook.active["F7"].value is None
            assert (".pdf" if format == "pdf" else ".xlsx") in result.headers["content-disposition"]
            app.dependency_overrides[get_current_user] = lambda: User(
                user_id="other-teacher", username="other", role="TEACHER", is_active=True)
            assert (await client.get(f"/api/reports/status/{task_id}")).status_code == 404
            assert (await client.get(f"/api/reports/download/{task_id}?format={format}")).status_code == 404
    finally:
        app.dependency_overrides.pop(get_current_user, None)
        if task_id:
            router_reports.TASKS.pop(task_id, None)


@pytest.mark.asyncio
async def test_empty_report_is_rejected_before_scheduling():
    app.dependency_overrides[get_current_user] = lambda: User(
        user_id="empty-teacher", username="empty", role="TEACHER", is_active=True)
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.post("/api/reports/generate", json={
                "report_kind": "gradebook", "format": "pdf", "custom_records": [],
            })
            assert response.status_code == 422
            assert "нет записей" in response.json()["detail"]
    finally:
        app.dependency_overrides.pop(get_current_user, None)


@pytest.mark.asyncio
async def test_pdf_renderer_failure_is_failed_instead_of_ready(monkeypatch, tmp_path):
    monkeypatch.setenv("REPORT_DIR", str(tmp_path))
    def broken_renderer(*args):
        raise RuntimeError("PDF unavailable")
    monkeypatch.setattr(router_reports, "generate_pdf_from_html", broken_renderer)
    task_id = str(uuid.uuid4())
    router_reports.TASKS[task_id] = {"status": "processing"}
    try:
        await router_reports._background_generate_report(task_id, custom_records=RECORDS,
                                                         report_kind="gradebook", output_format="pdf")
        assert router_reports.TASKS[task_id]["status"] == "failed"
        assert router_reports.TASKS[task_id]["error"]
        assert not list(tmp_path.iterdir())
    finally:
        router_reports.TASKS.pop(task_id, None)
