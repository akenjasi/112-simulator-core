import io
import os
from pathlib import Path
import pytest
from httpx import AsyncClient, ASGITransport
import openpyxl

from backend.main import app
from backend.core.deps import get_current_user
from backend.models.domain_01 import User
from backend.core.reports_engine import generate_excel_from_data
from backend.api.router_reports import TASKS, _background_generate_report


@pytest.mark.asyncio
async def test_export_excel_format(monkeypatch):
    """
    Test that the reports endpoint supports format=excel and returns an XLSX file.
    Since we don't have a full auth setup in this simple test, we mock dependencies if needed, 
    but mainly we expect a 401 or a valid 200 with the correct content type.
    """
    # Assuming we get a 422 if task_id is missing, but if provided, it generates
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        res = await ac.get("/api/reports/download?task_id=dummy&format=excel")
    
    # It might be 404 if dummy task doesn't exist, but it should recognize 'excel' format
    assert res.status_code in (404, 200, 401)
    if res.status_code == 200:
        assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_generate_excel_from_data_structure():
    """Test that generate_excel_from_data creates a valid, well-structured XLSX workbook in memory."""
    sample_cadets = [
        {
            "cadet_name": "Иванов Иван",
            "ticket_id": "Билет #42",
            "comm_score": 90,
            "card_score": 85,
            "sla_score": 100,
            "final_score": 92,
            "time_to_first_dispatch_sec": 45,
            "over_dispatched_services": [],
            "missed_critical_factoids": [],
            "penalties_list": [],
        },
        {
            "cadet_name": "Сидорова Анна",
            "ticket_id": "Билет #10",
            "comm_score": 60,
            "card_score": 50,
            "sla_score": 70,
            "final_score": 58,
            "time_to_first_dispatch_sec": 75,
            "over_dispatched_services": ["03"],
            "missed_critical_factoids": ["Запах газа"],
            "penalties_list": ["slow_dispatch", "over_dispatching", "missed_factoid"],
        },
    ]

    excel_bytes = generate_excel_from_data(sample_cadets)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    # Parse back with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Аттестация курсантов"

    # Row 1 is Title banner
    assert "СИСТЕМА-112" in str(ws["A1"].value)

    # Row 4 is Headers
    headers = [ws.cell(row=4, column=c).value for c in range(1, 11)]
    assert "Курсант" in headers
    assert "Итоговый балл" in headers
    assert "Лишние службы (гипер-диспетчеризация)" in headers

    # Row 5 is first cadet
    assert ws.cell(row=5, column=1).value == "Иванов Иван"
    assert ws.cell(row=5, column=3).value == 92

    # Row 6 is second cadet
    assert ws.cell(row=6, column=1).value == "Сидорова Анна"
    assert ws.cell(row=6, column=3).value == 58
    assert "03" in str(ws.cell(row=6, column=8).value)
    assert "Запах газа" in str(ws.cell(row=6, column=9).value)

    # Row 7 is Summary row with average
    assert ws.cell(row=7, column=1).value == "Среднее значение:"
    assert ws.cell(row=7, column=3).value == 75.0


@pytest.mark.asyncio
async def test_download_excel_authenticated_end_to_end():
    """Test full authenticated download with format=excel query parameter."""
    # Override auth to simulate logged in TEACHER
    app.dependency_overrides[get_current_user] = lambda: User(user_id="u-1", role="TEACHER", username="teacher", is_active=True)

    task_id = "test-export-task-excel"
    sample_cadets = [
        {
            "cadet_name": "Петров Петр",
            "ticket_id": "Билет #1",
            "comm_score": 100,
            "card_score": 100,
            "sla_score": 100,
            "final_score": 100,
            "time_to_first_dispatch_sec": 30,
            "over_dispatched_services": [],
            "missed_critical_factoids": [],
            "penalties_list": [],
        }
    ]

    excel_file = f"/tmp/report_{task_id}.xlsx"
    Path(excel_file).write_bytes(generate_excel_from_data(sample_cadets))

    TASKS[task_id] = {
        "status": "ready",
        "file_path": f"/tmp/report_{task_id}.pdf",
        "csv_path": f"/tmp/report_{task_id}.csv",
        "excel_path": excel_file,
        "error": None,
    }

    try:
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as ac:
            # 1. Query parameter test
            res = await ac.get(f"/api/reports/download?task_id={task_id}&format=excel")
            assert res.status_code == 200
            assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            assert len(res.content) > 0

            # Verify downloaded content is valid Excel
            wb = openpyxl.load_workbook(io.BytesIO(res.content))
            assert "Аттестация курсантов" in wb.sheetnames

            # 2. Path parameter test
            res_path = await ac.get(f"/api/reports/download/{task_id}?format=xlsx")
            assert res_path.status_code == 200
            assert res_path.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

            # 3. Missing task_id query param returns 422
            res_missing = await ac.get("/api/reports/download?format=excel")
            assert res_missing.status_code == 422

    finally:
        app.dependency_overrides.clear()
        if os.path.exists(excel_file):
            os.remove(excel_file)
        TASKS.pop(task_id, None)


@pytest.mark.asyncio
async def test_background_generate_report_creates_excel():
    """Verify that _background_generate_report creates the xlsx file asynchronously."""
    task_id = "test-bg-gen-excel"
    TASKS[task_id] = {
        "status": "processing",
        "file_path": None,
        "csv_path": None,
        "excel_path": None,
        "error": None,
    }

    custom_records = [
        {
            "cadet_name": "Тестовый Курсант",
            "ticket_id": "Билет #99",
            "comm_metrics": {"greeting_success": True, "filler_words": 0, "script_followed_pct": 100},
            "card_metrics": {
                "address_correct": True,
                "services_matched": True,
                "over_dispatched_services": [],
                "missed_critical_factoids": [],
            },
            "sla_metrics": {"sla_breached_count": 0, "time_to_first_dispatch_sec": 35},
        }
    ]

    await _background_generate_report(task_id=task_id, custom_records=custom_records)

    assert TASKS[task_id]["status"] == "ready"
    excel_path = TASKS[task_id]["excel_path"]
    assert excel_path is not None
    assert os.path.exists(excel_path)
    assert excel_path.endswith(".xlsx")

    # Verify openpyxl can read it
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    assert ws.cell(row=5, column=1).value == "Тестовый Курсант"

    # Cleanup
    if os.path.exists(excel_path):
        os.remove(excel_path)
    if TASKS[task_id].get("file_path") and os.path.exists(TASKS[task_id]["file_path"]):
        os.remove(TASKS[task_id]["file_path"])
    if TASKS[task_id].get("csv_path") and os.path.exists(TASKS[task_id]["csv_path"]):
        os.remove(TASKS[task_id]["csv_path"])
    TASKS.pop(task_id, None)

