import io
import openpyxl
import pytest
from sqlalchemy import select

from backend.database import AsyncSessionLocal
from backend.models.domain_02 import GeneratedTicket, ScenarioTicket


@pytest.mark.asyncio
async def test_batch_import_tickets(auth_client):
    """
    Test that template download works, and uploading a parsed list creates tickets.
    """
    # 1. Download template
    res_template = await auth_client.get("/api/tickets/import/template")
    assert res_template.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res_template.headers["content-type"]
    assert "tickets_template.xlsx" in res_template.headers.get("content-disposition", "")

    # Check template structure
    wb = openpyxl.load_workbook(io.BytesIO(res_template.content))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    expected_headers = [
        "Название",
        "Описание",
        "Номер звонящего",
        "Текст абонента",
        "Целевая служба (01, 02)",
        "Обязательные фактоиды (через запятую)",
    ]
    assert headers == expected_headers

    # 2. Upload excel file with tickets
    wb_new = openpyxl.Workbook()
    ws_new = wb_new.active
    ws_new.append(expected_headers)
    ws_new.append([
        "Пожар на балконе",
        "Горит балкон на 5 этаже",
        "+79001234567",
        "Срочно помогите, горит балкон, огонь поднимается наверх!",
        "01",
        "адрес, этаж, пламя",
    ])
    ws_new.append([
        "ДТП со скорой",
        "Столкновение на перекрестке с пострадавшим",
        "+79007654321",
        "Авария на перекрестке, водитель без сознания!",
        "01, 02, 03",
        "адрес, пострадавшие, скорая",
    ])
    stream = io.BytesIO()
    wb_new.save(stream)
    stream.seek(0)

    res_upload = await auth_client.post(
        "/api/tickets/import/upload",
        files={"file": ("tickets.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res_upload.status_code == 200
    data = res_upload.json()
    assert data["success"] is True
    assert data["success_count"] == 2
    assert data["error_count"] == 0

    # 3. Verify tickets exist in database
    async with AsyncSessionLocal() as session:
        stmt = select(ScenarioTicket).where(ScenarioTicket.title.in_(["Пожар на балконе", "ДТП со скорой"]))
        result = await session.execute(stmt)
        tickets = result.scalars().all()
        assert len(tickets) == 2

        fire_ticket = next(t for t in tickets if t.title == "Пожар на балконе")
        assert fire_ticket.category == "Пожары и задымления"
        assert fire_ticket.settings.get("etalon_services") == ["01 Пожарные"]

        dtp_ticket = next(t for t in tickets if t.title == "ДТП со скорой")
        assert "01 Пожарные" in dtp_ticket.settings.get("etalon_services")
        assert "02 Полиция" in dtp_ticket.settings.get("etalon_services")
        assert "03 Скорая" in dtp_ticket.settings.get("etalon_services")


@pytest.mark.asyncio
async def test_download_template_v1_router(auth_client):
    """Verify template download endpoint works on /api/v1/tickets/import/template as well."""
    res = await auth_client.get("/api/v1/tickets/import/template")
    assert res.status_code == 200
    assert len(res.content) > 0


@pytest.mark.asyncio
async def test_upload_tickets_validation_errors(auth_client):
    """
    Test validation errors during upload:
    Row 2: Valid
    Row 3: Valid
    Row 4: Empty target service ('строка 4 пустая целевая служба')
    Row 5: Empty title and caller text
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "Название",
        "Описание",
        "Номер звонящего",
        "Текст абонента",
        "Целевая служба (01, 02)",
        "Обязательные фактоиды (через запятую)",
    ])
    # Row 2 (valid)
    ws.append(["Утечка газа", "Запах газа", "+79110001122", "В подъезде сильно пахнет газом!", "04", "газ, подъезд"])
    # Row 3 (valid)
    ws.append(["Драка во дворе", "Шумят под окнами", "+79112223344", "Драка во дворе, вызовите полицию!", "02", "двор, драка"])
    # Row 4 (empty target service)
    ws.append(["Неизвестное происшествие", "Что-то происходит", "+79113334455", "Помогите нам пожалуйста!", "", "помощь"])
    # Row 5 (empty title and caller text)
    ws.append(["", "Только описание", "+79114445566", "", "01", "фактоид"])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = await auth_client.post(
        "/api/tickets/import/upload",
        files={"file": ("batch.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["success_count"] == 2
    assert data["error_count"] == 2
    assert any("строка 4 пустая целевая служба" in err.lower() for err in data["errors"])
    assert any("строка 5 пустое название или текст абонента" in err.lower() for err in data["errors"])


@pytest.mark.asyncio
async def test_upload_invalid_file_format_and_empty(auth_client):
    """Test rejection of non-excel files and empty content."""
    # 1. Non-excel extension
    res = await auth_client.post(
        "/api/tickets/import/upload",
        files={"file": ("tickets.txt", b"plain text", "text/plain")},
    )
    assert res.status_code == 400

    # 2. Empty file
    res = await auth_client.post(
        "/api/tickets/import/upload",
        files={"file": ("empty.xlsx", b"", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 400

    # 3. Corrupted excel file
    res = await auth_client.post(
        "/api/tickets/import/upload",
        files={"file": ("corrupt.xlsx", b"invalid zip content", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 400
